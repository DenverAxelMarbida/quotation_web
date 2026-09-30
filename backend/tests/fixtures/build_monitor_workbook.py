"""Builders for monitoring workbooks used by the import tests.

Two sources are used on purpose:

- ``generated_monitoring_workbook`` runs the real Excel generator, so at least one
  test proves that a workbook this application actually produces is readable by
  the import service. That is the round trip the whole feature depends on.
- The remaining helpers hand-build workbooks that the generator never produces,
  such as a file with the wrong columns or one saved as plain bytes.

No real business data is used: every client, project and description here is
invented (AGENTS.md sections 10 and 13).
"""

from io import BytesIO

from openpyxl import Workbook

from app.excel.openpyxl_generator import OpenpyxlWorkbookGenerator
from app.models.monitor import ImportedMonitorRow
from app.models.quotation import Quotation, QuotationItem
from app.models.workbook import ConfirmedQuotation, ConsolidatedWorkbookRequest

# The exact Summary header the generator writes, in order. Import must find these
# nine names, so the fixtures use the real list rather than a shortened copy.
SUMMARY_COLUMNS = list(OpenpyxlWorkbookGenerator.COLUMNS)


def _quotation(
    client: str,
    project: str,
    items: list[tuple[str, str, float | None, str | None]],
) -> Quotation:
    return Quotation(
        quotation_number=None,
        client_name=client,
        project_name=project,
        items=[
            QuotationItem(sr=sr, description=description, quantity=quantity, unit=unit)
            for sr, description, quantity, unit in items
        ],
    )


def generated_monitoring_workbook() -> bytes:
    """A workbook produced by the real generator, as the user would download it.

    Two sequence numbers and three rows, so a test can prove that grouping and
    row order survive the trip back in.
    """
    request = ConsolidatedWorkbookRequest(
        quotations=[
            ConfirmedQuotation(
                sequence_number="001",
                quotation=_quotation(
                    "SAMPLE CLIENT TRADING L.L.C",
                    "MBRC 466",
                    [
                        ("1.1", "Sample flooring product", 34.0, "m2"),
                        ("1.130", "Sample threshold", 88.0, "L.M."),
                    ],
                ),
            ),
            ConfirmedQuotation(
                sequence_number="002",
                quotation=_quotation(
                    "SAMPLE CLIENT TRADING L.L.C",
                    "MARINA BAY TOWER",
                    [("2.42", "Self-levelling compound", 122.0, "m2")],
                ),
            ),
        ]
    )
    return OpenpyxlWorkbookGenerator().generate_consolidated(request)


def combined_monitoring_workbook() -> bytes:
    """A workbook built from an existing monitoring file plus a new quotation.

    The situation this reproduces is the one the feature exists for. The user
    opens the monitoring workbook they already share, which by now has a
    schedule, a date and a status filled in by hand, and then adds a new
    quotation to that same file. Generating it has to produce that file with the
    new quotation underneath, not a replacement holding only the quotation.

    The existing rows are written in an order that is not numeric -- 007 before
    003 -- so a round trip that quietly sorted them would not pass.
    """
    request = ConsolidatedWorkbookRequest(
        existing_rows=[
            ImportedMonitorRow(
                sequence_number="007",
                client_name="SAMPLE CLIENT TRADING L.L.C",
                project_name="MARINA BAY TOWER",
                product_description="Sample skirting profile",
                quantity=18.0,
                unit_of_measurement="L.M.",
                installation_schedule="15-20 Nov 2026",
                start_date="2026-11-15",
                status="Ongoing",
            ),
            ImportedMonitorRow(
                sequence_number="003",
                client_name="SAMPLE CLIENT TRADING L.L.C",
                project_name="MBRC 466",
                product_description="Sample threshold strip",
                quantity=6.0,
                unit_of_measurement="L.M.",
                installation_schedule="",
                start_date=None,
                status="On Hold",
            ),
        ],
        quotations=[
            ConfirmedQuotation(
                sequence_number="008",
                quotation=_quotation(
                    "SAMPLE CLIENT TRADING L.L.C",
                    "CREEK VILLA",
                    [("1.1", "Sample engineered oak flooring", 34.0, "m2")],
                ),
            )
        ],
    )
    return OpenpyxlWorkbookGenerator().generate_consolidated(request)


def existing_monitor_rows() -> list[ImportedMonitorRow]:
    """Two rows as a monitoring workbook holds them, for request-model tests.

    Deliberately out of numeric order (007 before 003) and with a leading zero, so
    a test that cares about order or about sequences being identifiers can tell
    the difference. Every field is stated, because these tests are about what
    survives the request rather than about any default.

    No real business data: invented names throughout (AGENTS.md sections 10
    and 13).
    """
    return [
        ImportedMonitorRow(
            sequence_number="007",
            client_name="SAMPLE CLIENT TRADING L.L.C",
            project_name="MARINA BAY TOWER",
            product_description="Sample skirting profile",
            quantity=18.0,
            unit_of_measurement="L.M.",
            installation_schedule="15-20 Nov 2026",
            start_date="2026-11-15",
            status="Ongoing",
        ),
        ImportedMonitorRow(
            sequence_number="003",
            client_name="SAMPLE CLIENT TRADING L.L.C",
            project_name="MBRC 466",
            product_description="Sample threshold strip",
            quantity=6.0,
            unit_of_measurement="L.M.",
            installation_schedule="",
            start_date=None,
            status="On Hold",
        ),
    ]


def summary_workbook(
    rows: list[dict[str, object]],
    columns: list[str] | None = None,
    sheet_name: str = "Summary",
    extra_sheets: list[str] | None = None,
) -> bytes:
    """A workbook whose Summary sheet holds exactly the given rows.

    Column order comes from ``columns`` (the real nine by default) and each row is
    a mapping of column name to value, so a test only has to state the cells it
    cares about; anything it leaves out is written as an empty cell.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    header = SUMMARY_COLUMNS if columns is None else columns

    for index, name in enumerate(header, start=1):
        sheet.cell(row=1, column=index, value=name)

    for offset, row in enumerate(rows, start=2):
        for index, name in enumerate(header, start=1):
            sheet.cell(row=offset, column=index, value=row.get(name))

    for name in extra_sheets or []:
        workbook.create_sheet(title=name)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def row(
    sequence_number: object = "001",
    client_name: object = "SAMPLE CLIENT TRADING L.L.C",
    project_name: object = "MBRC 466",
    product_description: object = "Sample flooring product",
    quantity: object = 34.0,
    unit_of_measurement: object = "m2",
    installation_schedule: object = None,
    start_date: object = None,
    status: object = None,
) -> dict[str, object]:
    """One Summary row. Named arguments mirror the column names."""
    return {
        "Sequence Number": sequence_number,
        "Client Name": client_name,
        "Project Name": project_name,
        "Product Description": product_description,
        "Quantity": quantity,
        "Unit of Measurement": unit_of_measurement,
        "Installation Schedule": installation_schedule,
        "Start Date": start_date,
        "Status": status,
    }


def damaged_workbook_bytes() -> bytes:
    """Bytes that are not a readable workbook at all."""
    return b"this is not a spreadsheet, it is a sentence"


def truncated_workbook_bytes() -> bytes:
    """A real workbook whose container has been cut short.

    The file starts with the .xlsx zip signature, so only opening it reveals that
    the archive is incomplete.
    """
    complete = summary_workbook([row()])
    return complete[: len(complete) // 2]


__all__ = [
    "SUMMARY_COLUMNS",
    "combined_monitoring_workbook",
    "damaged_workbook_bytes",
    "existing_monitor_rows",
    "generated_monitoring_workbook",
    "row",
    "summary_workbook",
    "truncated_workbook_bytes",
]
