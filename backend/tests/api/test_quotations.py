"""API tests.

The extraction service is replaced through FastAPI dependency overrides, so
these tests exercise routing, validation and the error envelope without needing
a real PDF.
"""

import pytest
from fastapi.testclient import TestClient

from app.dependencies import get_extraction_service
from app.main import app
from app.models.preview import QuotationPreview, SourceInfo
from app.models.quotation import Quotation, QuotationItem
from app.models.review import ReviewFlag, ReviewReason
from app.services.errors import EmptyDocumentError
from tests.fixtures.build_monitor_workbook import existing_monitor_rows

client = TestClient(app)


def preview() -> QuotationPreview:
    return QuotationPreview(
        quotation=Quotation(
            quotation_number="Q-1",
            client_name="ACME",
            project_name="P-1",
            items=[QuotationItem(sr="1", description="flooring", quantity=30, unit="m2")],
        ),
        review=[
            ReviewFlag(
                field="items[0].description",
                reason=ReviewReason.UNCONFIDENT,
                message="please check",
            )
        ],
        source=SourceInfo(filename="quote.pdf", page_count=1),
    )


class StubService:
    def build_preview(self, pdf_bytes: bytes, filename: str) -> QuotationPreview:
        if not pdf_bytes.startswith(b"%PDF"):
            raise EmptyDocumentError("No text could be read from this PDF.")
        return preview()


@pytest.fixture(autouse=True)
def override_service() -> None:
    app.dependency_overrides[get_extraction_service] = lambda: StubService()
    yield
    app.dependency_overrides.clear()


def test_health_is_available() -> None:
    response = client.get("/api/health")

    assert response.status_code == 200
    assert response.json()["status"] == "ok"


def test_parse_returns_structured_data_with_review_flags() -> None:
    response = client.post(
        "/api/quotations/parse",
        files={"file": ("quote.pdf", b"%PDF-1.4 fake", "application/pdf")},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["quotation"]["quotation_number"] == "Q-1"
    assert body["quotation"]["items"][0]["quantity"] == 30
    assert body["review"][0]["field"] == "items[0].description"
    assert body["source"]["page_count"] == 1


def test_parse_requires_a_file() -> None:
    response = client.post("/api/quotations/parse")

    assert response.status_code == 422


def test_extraction_failures_return_the_structured_error_envelope() -> None:
    response = client.post(
        "/api/quotations/parse",
        files={"file": ("quote.pdf", b"not a pdf", "application/pdf")},
    )

    assert response.status_code == 422
    error = response.json()["error"]
    assert error["code"] == "EmptyDocumentError"
    assert "PDF" in error["message"]


def test_openapi_schema_is_served() -> None:
    response = client.get("/openapi.json")

    assert response.status_code == 200
    paths = response.json()["paths"]
    assert "/api/health" in paths
    assert "/api/quotations/parse" in paths
    assert "/api/quotations/generate-consolidated-excel" in paths


def _confirmed(seq: str, client_name: str, items: list[tuple[int, str, float, str]]):
    return {
        "sequence_number": seq,
        "quotation": {
            "quotation_number": "Q-1",
            "client_name": client_name,
            "project_name": "P-1",
            "items": [
                # The SR# travels as printed text, the way the parser emits it.
                {"sr": str(sr), "description": desc, "quantity": qty, "unit": unit}
                for sr, desc, qty, unit in items
            ],
        },
    }


def test_generate_consolidated_returns_one_downloadable_workbook() -> None:
    from io import BytesIO

    from openpyxl import load_workbook

    payload = {
        "quotations": [
            _confirmed("001", "ELEV8", [(1, "A", 153.0, "m2"), (2, "B", 18.0, "STEP")]),
            _confirmed("002", "TEE VEE EFF FZCO", [(1, "C", 46.0, "m2")]),
        ]
    }

    response = client.post("/api/quotations/generate-consolidated-excel", json=payload)

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment" in response.headers["content-disposition"]

    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook["Summary"]
    sequence_column = [sheet.cell(row=r, column=1).value for r in range(2, sheet.max_row + 1)]
    assert sequence_column == ["001", "001", "002"]


def test_generate_consolidated_rejects_a_request_with_nothing_to_write() -> None:
    # Both lists empty is the one combination that is not a job. The generator
    # would answer it with a header-only file, which the user would open
    # expecting a monitoring sheet and find an empty one.
    response = client.post("/api/quotations/generate-consolidated-excel", json={"quotations": []})

    assert response.status_code == 422


def test_generate_consolidated_rejects_a_request_naming_no_source_at_all() -> None:
    # The same mistake phrased differently. The two fields are individually
    # optional now, so this is what an unbuilt request body actually looks like.
    response = client.post("/api/quotations/generate-consolidated-excel", json={})

    assert response.status_code == 422


def test_the_empty_request_error_says_what_to_provide() -> None:
    # A 422 the user cannot act on is the same as no answer. The two possible
    # sources have to be named, or she still does not know which one to give.
    response = client.post("/api/quotations/generate-consolidated-excel", json={"quotations": []})

    detail = str(response.json()["detail"])
    assert "quotations" in detail
    assert "existing_rows" in detail


def test_generate_consolidated_exports_an_imported_monitor_on_its_own() -> None:
    # The monitor-only workflow: she opened the workbook, corrected a row and
    # wants the corrected copy back. No quotation is involved, and requiring one
    # would mean inventing work to get a file out of the system.
    from io import BytesIO

    from openpyxl import load_workbook

    rows = existing_monitor_rows()
    response = client.post(
        "/api/quotations/generate-consolidated-excel",
        json={"quotations": [], "existing_rows": [row.model_dump(mode="json") for row in rows]},
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    sheet = load_workbook(BytesIO(response.content))["Summary"]
    assert [sheet.cell(row=r, column=1).value for r in range(2, sheet.max_row + 1)] == [
        "007",
        "003",
    ]


def test_a_monitor_only_export_keeps_the_rows_as_they_stand() -> None:
    # The whole point is that the file comes back as she left it, including the
    # three fields only she controls and the sequence numbers she has been using
    # all year.
    from io import BytesIO

    from openpyxl import load_workbook

    rows = existing_monitor_rows()
    response = client.post(
        "/api/quotations/generate-consolidated-excel",
        json={"quotations": [], "existing_rows": [row.model_dump(mode="json") for row in rows]},
    )

    sheet = load_workbook(BytesIO(response.content))["Summary"]
    headers = [sheet.cell(row=1, column=c).value for c in range(1, sheet.max_column + 1)]
    first = [sheet.cell(row=2, column=c).value for c in range(1, sheet.max_column + 1)]

    assert headers[:9] == [
        "Sequence Number",
        "Client Name",
        "Project Name",
        "Product Description",
        "Quantity",
        "Unit of Measurement",
        "Installation Schedule",
        "Start Date",
        "Status",
    ]
    assert first == [
        "007",
        "SAMPLE CLIENT TRADING L.L.C",
        "MARINA BAY TOWER",
        "Sample skirting profile",
        18.0,
        "L.M.",
        "15-20 Nov 2026",
        "2026-11-15",
        "Ongoing",
    ]


def test_generate_consolidated_merges_an_imported_monitor_with_new_quotations() -> None:
    # Existing rows first and in their own order, then the new quotation, so the
    # sequences she already handed out keep their meaning.
    from io import BytesIO

    from openpyxl import load_workbook

    rows = existing_monitor_rows()
    response = client.post(
        "/api/quotations/generate-consolidated-excel",
        json={
            "quotations": [_confirmed("008", "ELEV8", [(1, "new line", 12.0, "m2")])],
            "existing_rows": [row.model_dump(mode="json") for row in rows],
        },
    )

    assert response.status_code == 200

    sheet = load_workbook(BytesIO(response.content))["Summary"]
    assert [sheet.cell(row=r, column=1).value for r in range(2, sheet.max_row + 1)] == [
        "007",
        "003",
        "008",
    ]


def test_a_monitor_only_export_is_never_a_header_only_file() -> None:
    # The degenerate case the request validator exists to prevent, asserted at the
    # level the user reaches it: whatever comes back has rows in it.
    from io import BytesIO

    from openpyxl import load_workbook

    rows = existing_monitor_rows()
    response = client.post(
        "/api/quotations/generate-consolidated-excel",
        json={"quotations": [], "existing_rows": [row.model_dump(mode="json") for row in rows]},
    )

    assert load_workbook(BytesIO(response.content))["Summary"].max_row > 1


def test_generate_consolidated_requires_sequence_number() -> None:
    response = client.post(
        "/api/quotations/generate-consolidated-excel",
        json={"quotations": [{"sequence_number": "", "quotation": {"items": []}}]},
    )

    assert response.status_code == 422
