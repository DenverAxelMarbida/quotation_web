"""Reading a monitoring workbook back in.

The Excel file is the work state that outlives the web application: it is
downloaded, kept in OneDrive, and opened again weeks later. This service is the
only place that knows how to read one, and it is deliberately strict, because
importing the wrong file would show the user a monitor that is not theirs.

The rules it enforces:

- The file is really a workbook, and it really has this application's Summary
  sheet with its nine columns. Anything else is refused by name.
- A Sequence Number is read as the identifier it is. ``"007"`` is never turned
  into ``7``; only the highest-sequence calculation compares numerically.
- A cell the user left empty stays empty. Blank Installation Schedule, Start
  Date and Status columns are normal, not errors, and none of them is filled in
  with a guess.
- A Status outside the dropdown's three values is reported. The column has a
  dropdown, so a fourth value means the file came from somewhere else.

AGENTS.md section 9: when something cannot be read, say so in words the user can
act on rather than presenting a half-filled sheet as if it were correct.
"""

from datetime import date, datetime
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.excel.openpyxl_generator import OpenpyxlWorkbookGenerator
from app.models.monitor import MONITOR_STATUS_VALUES, ImportedMonitor, ImportedMonitorRow
from app.services.errors import (
    EmptySummaryError,
    InvalidSequenceNumberError,
    InvalidStatusValueError,
    MissingColumnsError,
    MissingSummarySheetError,
    NotExcelError,
)

# Reusing the generator's own header list means the two cannot drift apart: if
# the workbook layout changes, the import check follows it.
REQUIRED_COLUMNS = tuple(OpenpyxlWorkbookGenerator.COLUMNS)
SUMMARY_SHEET = "Summary"

# The workbook stores dates as real dates and displays them in its own format, so
# that format is what gets read back rather than a different convention.
DATE_DISPLAY_FORMAT = "%d/%m/%Y"


class MonitorImportService:
    """Validate a monitoring workbook and return its rows."""

    def import_monitor(self, xlsx_bytes: bytes, filename: str) -> ImportedMonitor:
        """Read the Summary sheet of a monitoring workbook.

        Raises a :class:`MonitorImportError` subclass, each message written for
        the user, when the file cannot be imported.
        """
        rows = self._read_rows(xlsx_bytes)
        if not rows:
            raise EmptySummaryError(
                "The Summary sheet has no rows. Add a quotation to the workbook "
                "and save it before uploading it here."
            )

        highest = max(rows, key=lambda item: int(item.sequence_number))
        return ImportedMonitor(
            rows=rows,
            highest_sequence=highest.sequence_number,
            row_count=len(rows),
            source_filename=filename,
        )

    def _read_rows(self, xlsx_bytes: bytes) -> list[ImportedMonitorRow]:
        if not xlsx_bytes:
            raise NotExcelError("The uploaded file is empty.")

        try:
            workbook = load_workbook(BytesIO(xlsx_bytes), data_only=True, read_only=True)
        except Exception as error:
            # zipfile.BadZipFile, XML errors from a damaged file, and openpyxl's
            # own complaints all mean one thing to the user: not a readable file.
            raise NotExcelError(
                "The file could not be opened. Check that it is the .xlsx file "
                "produced by this application, and that it is not damaged."
            ) from error

        try:
            return self._read_summary_rows(workbook)
        except (InvalidFileException, KeyError, ValueError, TypeError) as error:
            raise NotExcelError(
                "The file could not be read as an Excel workbook. Check that it "
                "is the .xlsx file produced by this application."
            ) from error
        finally:
            workbook.close()

    def _read_summary_rows(self, workbook) -> list[ImportedMonitorRow]:
        if SUMMARY_SHEET not in workbook.sheetnames:
            raise MissingSummarySheetError(
                f"This file has no '{SUMMARY_SHEET}' sheet, so it is not a "
                "monitoring sheet. Upload the .xlsx file produced by this "
                "application."
            )

        sheet = workbook[SUMMARY_SHEET]
        header_index: dict[str, int] = {}
        rows: list[ImportedMonitorRow] = []

        for line, row in enumerate(sheet.iter_rows(), start=1):
            if all(cell.value in (None, "") for cell in row):
                # A row of empty cells is a spreadsheet artefact, not a quotation:
                # the generator leaves validation ranges below the data, and
                # Excel can save a trailing empty row.
                continue
            if not header_index:
                header_index = {
                    _column_name(cell.value): index
                    for index, cell in enumerate(row)
                    if cell.value not in (None, "")
                }
                self._require_columns(header_index)
                continue
            rows.append(self._read_row(row, header_index, line))

        return rows

    def _require_columns(self, header_index: dict[str, int]) -> None:
        present = set(header_index)
        missing = [name for name in REQUIRED_COLUMNS if _column_name(name) not in present]
        if not missing:
            return
        raise MissingColumnsError(
            f"The Summary sheet is missing these required columns: "
            f"{', '.join(missing)}. Upload the .xlsx file produced by this "
            "application, or restore the missing columns and try again."
        )

    def _read_row(self, row, header_index: dict[str, int], line: int) -> ImportedMonitorRow:
        def cell(name: str):
            index = header_index[_column_name(name)]
            return row[index].value if index < len(row) else None

        return ImportedMonitorRow(
            sequence_number=self._sequence_number(cell("Sequence Number"), line),
            client_name=_text(cell("Client Name")),
            project_name=_text(cell("Project Name")),
            product_description=_text(cell("Product Description")),
            quantity=_number(cell("Quantity")),
            unit_of_measurement=_text(cell("Unit of Measurement")),
            installation_schedule=_text(cell("Installation Schedule")),
            start_date=_optional_text(cell("Start Date")),
            status=self._status(cell("Status"), line),
        )

    def _sequence_number(self, value, line: int) -> str:
        if value in (None, ""):
            raise InvalidSequenceNumberError(
                f"Row {line} has no Sequence Number. Every row in a monitoring "
                "sheet must belong to a numbered quotation."
            )
        if isinstance(value, bool):
            raise InvalidSequenceNumberError(
                f"Row {line} has a Sequence Number this application cannot read. "
                "Use a whole number, for example 001."
            )
        if isinstance(value, (int, float)):
            if not float(value).is_integer():
                raise InvalidSequenceNumberError(
                    f"Row {line} has a Sequence Number this application cannot "
                    f"read: {value!r}. Use a whole number, for example 001."
                )
            # Excel displays an unformatted number without padding, so a
            # hand-typed 7 stays "7" rather than gaining zeros the file never had.
            return str(int(value))
        if isinstance(value, str) and value.strip().isdigit():
            # An identifier, not a number: "007" and "001" keep their zeros.
            return value.strip()
        raise InvalidSequenceNumberError(
            f"Row {line} has a Sequence Number this application cannot read: "
            f"{_text(value)!r}. Use a whole number, for example 001."
        )

    def _status(self, value, line: int) -> str:
        text = _optional_text(value)
        if text in (None, ""):
            return ""
        if text not in MONITOR_STATUS_VALUES:
            allowed = ", ".join(MONITOR_STATUS_VALUES)
            raise InvalidStatusValueError(
                f"Row {line} has the Status '{text}', which is not one of the "
                f"statuses this application uses ({allowed}). Leave the Status "
                "blank or choose one of those values in Excel, then upload the "
                "file again."
            )
        return text


def _column_name(value) -> str:
    return _text(value).strip().lower()


def _text(value) -> str:
    """A cell as readable text, with an empty cell becoming an empty string."""
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        # A number typed into a text column, e.g. a project called 466.
        return str(int(value))
    if isinstance(value, (datetime, date)):
        return value.strftime(DATE_DISPLAY_FORMAT)
    return str(value)


def _optional_text(value) -> str | None:
    """Like :func:`_text`, but an empty cell becomes None rather than ''."""
    text = _text(value)
    return text or None


def _number(value) -> float | None:
    """A numeric cell, or None when it is empty or not a number.

    A quantity that cannot be read is left empty instead of being replaced by a
    guess. The row is still imported so the user can see exactly which cell is
    missing a value.
    """
    if isinstance(value, bool) or value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        try:
            return float(value.strip())
        except ValueError:
            return None
    return None
